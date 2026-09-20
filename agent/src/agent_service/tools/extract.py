from __future__ import annotations

import re
from html.parser import HTMLParser
from io import BytesIO

_SKIP = {"script", "style", "nav", "noscript", "svg", "template"}
_BLOCK = {"p", "div", "section", "article", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr"}


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip = 0
        self._parts: list[str] = []
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        name = tag.lower()
        if name in _SKIP:
            self._skip += 1
            return
        if name == "title" and self._skip == 0:
            self._in_title = True
        if name in _BLOCK and self._skip == 0:
            self._parts.append("\n")
        if name in {"h1", "h2", "h3", "h4", "h5", "h6"} and self._skip == 0:
            self._parts.append("\n# ")

    def handle_endtag(self, tag: str) -> None:
        name = tag.lower()
        if name in _SKIP and self._skip:
            self._skip -= 1
            return
        if name == "title":
            self._in_title = False
        if name in _BLOCK and self._skip == 0:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        text = re.sub(r"\s+", " ", data).strip()
        if not text:
            return
        if self._in_title and not self.title:
            self.title = text
            return
        self._parts.append(text + " ")


def html_to_text(raw: str) -> tuple[str, str]:
    parser = _HTMLText()
    parser.feed(raw)
    parser.close()
    body = re.sub(r"\n{3,}", "\n\n", "".join(parser._parts)).strip()
    return parser.title, body


def pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(part for part in pages if part)
    if not text.strip():
        raise ValueError("unable to extract text from PDF (scanned or empty text layer)")
    return text


def docx_to_text(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    parts: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    text = "\n\n".join(parts)
    if not text.strip():
        raise ValueError("unable to extract text from DOCX")
    return text


def looks_like_pdf(content: bytes, content_type: str) -> bool:
    if content.startswith(b"%PDF"):
        return True
    return "pdf" in (content_type or "").lower()


def looks_like_html(content: bytes, content_type: str) -> bool:
    ctype = (content_type or "").lower()
    if "html" in ctype:
        return True
    head = content[:200].lstrip().lower()
    return head.startswith(b"<!doctype html") or head.startswith(b"<html")
